from docx import Document
from app.services.document_preview import parse_table,table_chunks
from docx.enum.text import WD_BREAK, WD_ALIGN_PARAGRAPH


def test_inline_breaks_and_tabs_keep_their_position_between_text_nodes():
    from app.services.document_preview import run_text
    doc = Document()
    run = doc.add_paragraph().add_run('First line')
    run.add_break()
    run.add_text('Second line')
    run.add_tab()
    run.add_text('Last field')
    run.add_break(WD_BREAK.PAGE)
    assert run_text(run._r) == 'First line\nSecond line\tLast field'


def test_export_keeps_alignment_and_multiline_table_content(tmp_path):
    from app.services.ocr_native_export import build_docx, build_html
    data = {'pages': [{'elements': [
        {'type': 'paragraph', 'text': 'Document title', 'align': 'center', 'is_heading': True},
        {'type': 'paragraph', 'text': 'Description', 'align': 'center'},
        {'type': 'paragraph', 'text': 'Body text', 'align': 'left'},
        {'type': 'table', 'xs': [0, 100], 'grid': [['A\nB\nC']], 'cells': [
            {'row': 0, 'col': 0, 'rowspan': 1, 'colspan': 1, 'text': 'A\nB\nC'}]}
    ]}]}
    path = tmp_path/'layout.docx'
    build_docx(data, path)
    doc = Document(path)
    assert [p.alignment for p in doc.paragraphs[:3]] == [
        WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT]
    assert doc.tables[0].cell(0, 0).text == 'A\nB\nC'
    assert 'text-align:center' in build_html(data)


def test_preview_preserves_style_borders_and_both_merge_directions():
    doc=Document()
    table=doc.add_table(rows=6,cols=4)
    table.style='Table Grid'
    table.cell(0,0).merge(table.cell(0,3)).text='Merged heading'
    table.cell(1,0).merge(table.cell(3,0)).text='Three rows'
    def paragraph(p, **kwargs):
        return {'content':p.text or '', 'align':'left'}
    block=parse_table(table._tbl,paragraph,doc.styles.element)
    assert block['has_borders']
    assert block['cell_props'][0][0]['colspan']==4
    assert block['cell_props'][1][0]['rowspan']==3
    assert block['cell_props'][2][0]['hidden']
    assert block['cell_props'][3][0]['hidden']
    chunks=list(table_chunks(block,80))
    assert sum(len(c['data']) for c in chunks)==6
    assert all(c['has_borders'] and c['col_widths']==block['col_widths'] for c in chunks)
    assert any(len(c['data'])>=3 and c['cell_props'][0][0]['rowspan']==3 for c in chunks)


def test_measured_small_rows_and_merged_headers_keep_source_heights():
    from docx.shared import Pt
    doc = Document()
    table = doc.add_table(rows=4, cols=2)
    for row, height in zip(table.rows, [8, 9, 7, 6]):
        row.height = Pt(height)
    table.cell(0, 0).merge(table.cell(2, 0)).text = 'A measured three-line heading'
    table.cell(3, 1).text = '123.45'
    block = parse_table(table._tbl, lambda p, **kw: {'content':p.text}, doc.styles.element)
    assert block['row_heights'] == [8, 9, 7, 6]
    assert block['cell_props'][0][0]['rowspan'] == 3


def test_table_fit_reduces_excess_leading_and_single_glyph_outliers():
    from docx.shared import Pt
    from app.services.ocr_native_export import fit_measured_table_rows
    doc = Document()
    table = doc.add_table(rows=2, cols=2)
    for row in table.rows:
        row.height = Pt(10)
    for cell in table._cells:
        p = cell.paragraphs[0]
        p.add_run('123').font.size = Pt(7)
        p.paragraph_format.line_spacing = Pt(8)
        p.paragraph_format.space_before = Pt(4)
    p = table.cell(0, 0).paragraphs[0]
    p.runs[0].text = 'x'
    p.runs[0].font.size = Pt(24)
    p.paragraph_format.line_spacing = Pt(27)
    texts = [c.text for c in table._cells]
    fit_measured_table_rows(table)
    assert [c.text for c in table._cells] == texts
    assert p.paragraph_format.line_spacing.pt <= 10
    assert p.runs[0].font.size.pt < 10
    other = table.cell(1, 1).paragraphs[0]
    assert other.runs[0].font.size.pt == 7
    assert other.paragraph_format.space_before.pt < 2
    assert all(row.height.pt == 10 for row in table.rows)


def test_table_fit_uses_whole_vertical_span_and_preserves_extra_content():
    from docx.shared import Pt
    from app.services.ocr_native_export import fit_measured_table_rows
    doc = Document()
    table = doc.add_table(rows=3, cols=2)
    for row in table.rows:
        row.height = Pt(10)
    merged = table.cell(0, 0).merge(table.cell(2, 0)).paragraphs[0]
    merged.add_run('Title').font.size = Pt(18)
    merged.paragraph_format.line_spacing = Pt(21)
    extra = table.cell(0, 1).paragraphs[0]
    extra.add_run('Many\nextra\nlines').font.size = Pt(12)
    extra.paragraph_format.line_spacing = Pt(14)
    fit_measured_table_rows(table)
    assert merged.runs[-1].font.size.pt == 18
    assert extra.runs[-1].font.size.pt == 12
    assert extra.text == 'Many\nextra\nlines'
