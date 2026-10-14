import base64
import io

import cv2
import fitz
import numpy as np
from PIL import Image
from docx import Document

from app.services.ocr_native import detect_tables, _annotate_alignment, _is_graphic_residue
from app.services.ocr_margins import repeated_footer_regions
from app.services.ocr_native_export import build_docx
from app.services.document_preview import run_images
from app.services.ocr_orientation import detect_rotation, rotate_image


def test_speckle_removal_keeps_stacked_accent_components():
    from app.services.ocr_native import _remove_distant_speckles
    binary=np.full((120,240),255,np.uint8)
    for x in [30,80,130,180]:
        cv2.rectangle(binary,(x,70),(x+12,99),0,-1)
    cv2.rectangle(binary,(33,62),(40,66),0,-1)
    cv2.rectangle(binary,(34,54),(39,57),0,-1)
    cv2.rectangle(binary,(210,5),(212,7),0,-1)
    cleaned=_remove_distant_speckles(binary.copy())
    assert np.all(cleaned[54:58,34:40]==0)
    assert np.all(cleaned[62:67,33:41]==0)
    assert np.all(cleaned[5:8,210:213]==255)


def test_justified_cell_words_remain_on_one_source_line():
    from app.services.ocr_native import _lines
    words = [dict(text=text,bbox=[x,20,x+30,40],confidence=96,line=(1,1,1))
             for text,x in [('Some',10),('widely',55),('spaced',110),('words',250)]]
    assert len(_lines(words,0)) == 2  # A page-wide header may have a gutter.
    lines = _lines(words,0,split_columns=False)
    assert len(lines)==1 and lines[0]['text']=='Some widely spaced words'


def test_empty_thin_cell_does_not_expand_to_default_font_height(tmp_path):
    table = dict(type='table',bbox=[0,0,400,110],xs=[0,200,400],ys=[0,10,110],
                 grid=[['',''],['Value','Other']],cells=[])
    for r,row in enumerate(table['grid']):
        for c,text in enumerate(row):
            table['cells'].append(dict(row=r,col=c,rowspan=1,colspan=1,text=text))
    path=tmp_path/'thin.docx'
    build_docx({'pages':[dict(page_w=500,page_h=700,elements=[table])]},path)
    paragraph=Document(path).tables[0].cell(0,0).paragraphs[0]
    assert paragraph.paragraph_format.line_spacing.pt==1
    assert paragraph.runs[0].font.size.pt==1


def test_marginal_image_keeps_page_position_without_body_height(tmp_path):
    image=Image.new('RGB',(100,30),'gray');buffer=io.BytesIO();image.save(buffer,format='PNG')
    path=tmp_path/'anchor.docx'
    page=dict(page_w=1000,page_h=1400,page_width_pt=500,page_height_pt=700,elements=[
        dict(type='paragraph',text='Body',bbox=[100,100,400,150]),
        dict(type='image',role='footer',bbox=[800,1320,900,1350],width_fraction=.1,
             data=base64.b64encode(buffer.getvalue()).decode('ascii'))])
    build_docx({'pages':[page]},path)
    doc=Document(path)
    assert len(doc.paragraphs)==1
    html,height=run_images(doc.paragraphs[0].runs[-1]._r,doc.part)
    assert height==0
    assert 'position:absolute;left:400.00pt;top:660.00pt' in html
    assert doc.sections[0].right_margin.pt==300


def test_single_row_continuation_keeps_every_column():
    gray = np.full((600,900),255,np.uint8)
    for x in [40,190,420,670,850]:
        cv2.line(gray,(x,40),(x,470),0,2)
    for y in [40,470]:
        cv2.line(gray,(40,y),(850,y),0,2)
    tables,_ = detect_tables(gray)
    assert len(tables)==1
    assert len(tables[0]['grid'])==1 and len(tables[0]['grid'][0])==4
    assert len(tables[0]['cells'])==4
    # Two decorative rules alone must not become a table.
    gray[:]=255
    for y in [40,470]:
        cv2.line(gray,(40,y),(850,y),0,2)
    assert detect_tables(gray)[0]==[]


def test_header_gutter_allows_unequal_columns_and_stops_at_body():
    from app.services.ocr_native import _header_columns, _lines
    words=[]
    for row,y in enumerate([50,85,125]):
        for col,positions in enumerate([[80,150,220],[420,510,590,670]]):
            if row==1 and col==0:
                continue
            for x in positions:
                words.append(dict(text='Example',bbox=[x,y,x+55,y+22],confidence=95,line=(1,row,1)))
    words.append(dict(text='An independent centered title',bbox=[260,200,680,224],confidence=95,line=(2,1,1)))
    lines=_lines(words,0)
    columns=_header_columns(lines,900,1200)
    assert columns is not None and [len(c) for c in columns]==[2,3]
    assert all(l['bbox'][2]<350 for l in columns[0])
    assert all(l['bbox'][0]>350 for l in columns[1])


def test_cleaning_cell_edges_keeps_leading_letters_and_detached_accents():
    from app.services.ocr_native import _prepare_cell
    gray=np.full((200,440),255,np.uint8)
    for y in [50,110,170]:
        cv2.putText(gray,'Sample row',(4,y),cv2.FONT_HERSHEY_SIMPLEX,1,0,2)
    # A detached accent over a glyph, plus unrelated debris at the far right.
    gray[21:24,11:17]=0
    gray[15:17,390:420]=0
    prepared=_prepare_cell(gray)
    assert prepared is not None
    clean=prepared[0]
    assert np.count_nonzero(clean[:,0:25]<128)>150
    assert np.count_nonzero(clean[:10,:30]<128)>0
    assert clean.shape[1]<350


def test_export_uses_source_dimensions_indents_and_line_breaks(tmp_path):
    from app.services.document_preview import page_setup
    path=tmp_path/'measured.docx'
    lines=[dict(text='Indented first line',bbox=[150,80,700,112],font_size_pt=14,bold=False),
           dict(text='A continued printed line',bbox=[100,128,680,160],font_size_pt=14,bold=False)]
    page=dict(page_w=1000,page_h=700,page_width_pt=720,page_height_pt=504,
              elements=[dict(type='paragraph',text='Joined text',bbox=[100,80,700,160],source_lines=lines)])
    build_docx({'pages':[page]},path)
    doc=Document(path);p=doc.paragraphs[0]
    assert p.text=='Indented first line\nA continued printed line'
    assert abs(p.paragraph_format.first_line_indent.pt-36)<.1
    assert doc.sections[0].page_width.pt==720
    assert page_setup(doc.sections[0])['height_pt']==504
    assert all(run.font.size.pt==14 for run in p.runs)


def test_typography_measures_size_and_weight_from_pixels():
    from PIL import ImageDraw,ImageFont
    from app.services.ocr_typography import annotate_lines,_fonts
    fonts=_fonts()
    assert fonts
    for bold in [False,True]:
        font=next(f for family,b,i,f in fonts if family=='Arial' and b==bold and not i)
        im=Image.new('L',(1100,140),255);draw=ImageDraw.Draw(im);words=[];x=8
        for text in ['Several','different','words']:
            bbox=draw.textbbox((x,0),text,font=font);draw.text((x,0),text,font=font,fill=0)
            words.append(dict(text=text,bbox=list(bbox),confidence=96));x=bbox[2]+25
        line=dict(text='Several different words',words=words)
        annotate_lines([line],np.array(im))
        assert abs(line['font_size_pt']-24)<1.5
        assert line['bold']==bold


def test_narrow_rows_and_columns_do_not_create_empty_morphology_inputs():
    gray = np.full((500,600),255,np.uint8)
    for x in [50,220,228,480]:
        cv2.line(gray,(x,50),(x,430),0,1)
    for y in [50,200,208,430]:
        cv2.line(gray,(50,y),(480,y),0,1)
    tables,_ = detect_tables(gray)
    assert len(tables) == 1
    assert len(tables[0]['grid']) == 3
    assert len(tables[0]['grid'][0]) == 3
    assert all(c['bbox'][2] > c['bbox'][0] and c['bbox'][3] > c['bbox'][1] for c in tables[0]['cells'])


def test_letter_stem_in_merged_caption_is_not_a_column_border():
    gray = np.full((500,600),255,np.uint8)
    for x in [50,450]:
        cv2.line(gray,(x,50),(x,400),0,1)
    cv2.line(gray,(200,110),(200,400),0,1)
    for y in [50,110,240,400]:
        cv2.line(gray,(50,y),(450,y),0,1)
    cv2.line(gray,(200,68),(200,88),0,2)
    tables,_ = detect_tables(gray)
    assert tables[0]['cells'][0]['colspan'] == 2


def test_export_preserves_mixed_portrait_and_landscape_pages(tmp_path):
    path = tmp_path/'orientation.docx'
    pages = [dict(page_w=w,page_h=h,elements=[dict(type='paragraph',text='Page')])
             for w,h in [(800,1100),(1100,800),(800,1100)]]
    build_docx({'pages':pages},path)
    doc = Document(path)
    assert [s.page_width>s.page_height for s in doc.sections] == [False,True,False]


def test_indented_paragraph_first_line_is_not_a_centered_title():
    lines = [dict(text='A paragraph starts here and continues,',bbox=[170,100,830,140],source_paragraph=('2','3')),
             dict(text='across the next printed line.',bbox=[100,165,900,205],source_paragraph=('2','3')),
             dict(text='A genuine title',bbox=[350,280,650,320],source_paragraph=('3','1'))]
    _annotate_alignment(lines,(100,900))
    assert [line['align'] for line in lines] == ['left','left','center']


def test_stamp_edge_fragments_require_pixel_evidence_not_just_low_confidence():
    original = np.full((80,100),255,np.uint8)
    original[20:60,20:80] = 80
    cleaned = np.full_like(original,255)
    cleaned[30:32,25:50] = 0
    line = dict(bbox=[20,20,80,60],confidence=40)
    assert _is_graphic_residue(line,original,cleaned)
    assert not _is_graphic_residue(line,cleaned,cleaned)
    cleaned[25:50,30:60] = 0
    assert not _is_graphic_residue(line,original,cleaned)


def test_repeated_margin_graphic_is_separated_but_unique_footnote_is_not():
    doc = fitz.open()
    for i in range(4):
        page = doc.new_page(width=595,height=842)
        page.insert_text((60,100),f'Different body text {i}')
        page.draw_rect(fitz.Rect(480,807,580,832))
        page.insert_text((486,824),'MARGIN LOGO',fontsize=9)
        if i == 1:
            page.insert_text((60,795),'Unique footnote stays readable',fontsize=8)
    regions = repeated_footer_regions(doc.tobytes())
    assert set(regions) == {0,1,2,3}
    assert all(len(boxes)==1 and boxes[0][0] > .75 for boxes in regions.values())


def test_marginal_image_survives_docx_export_and_preview(tmp_path):
    png = io.BytesIO()
    Image.new('RGB',(150,45),(80,100,200)).save(png,format='PNG')
    graphic = dict(type='image',role='footer',data=base64.b64encode(png.getvalue()).decode(),width_fraction=.2,align='right')
    path = tmp_path/'footer.docx'
    build_docx({'pages':[{'elements':[dict(type='paragraph',text='Body text'),graphic]}]},path)
    doc = Document(path)
    assert len(doc.inline_shapes) == 1
    assert doc.paragraphs[0].text == 'Body text'
    html,height = run_images(doc.paragraphs[-1].runs[0]._r,doc.part)
    assert 'data:image/png;base64,' in html
    assert height > 0


def test_orientation_probe_reads_sideways_text_without_rotating_an_empty_page():
    from app.services.ocr_native import _tesseract_words
    gray = np.full((1200,900),255,np.uint8)
    for i in range(18):
        cv2.putText(gray,'Printed document with several lines of text',
                    (30,65+i*58),cv2.FONT_HERSHEY_SIMPLEX,.9,0,2)
    sideways = rotate_image(gray,270)
    read = lambda probe:_tesseract_words(probe,'en',6,lambda:False)
    assert detect_rotation(sideways,lambda:False,read) == 90
    assert detect_rotation(np.full_like(gray,255),lambda:False,read) == 0
    assert np.array_equal(rotate_image(sideways,90),gray)
