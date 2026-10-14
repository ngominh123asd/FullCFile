from copy import deepcopy

import cv2
import numpy as np
from PIL import Image, ImageDraw
from docx import Document

from app.services.ocr_typography import _fonts, annotate_lines, harmonize_typography
from app.services.ocr_inline import script_candidates, recover_inline_scripts
from app.services.ocr_text_refinement import visual_accent_support
from app.services.ocr_text_refinement import refine_with_context
from app.services.ocr_region_refinement import refine_weak_spans
from app.services.ocr_native_export import build_docx


def font_at(size, bold=False, italic=False):
    return next(f for name,b,i,f in _fonts()
                if name=='Times New Roman' and (b,i)==(bold,italic)).font_variant(size=size)


def test_point_size_does_not_depend_on_spaces_missed_by_ocr():
    image=Image.new('L',(700,100),255);draw=ImageDraw.Draw(image);font=font_at(50)
    lines=[]
    for x,printed,recognized in [(10,'28 - 37','28-37'),(350,'123.456','123.456')]:
        box=list(draw.textbbox((x,10),printed,font=font));draw.text((x,10),printed,font=font,fill=0)
        lines.append(dict(text=recognized,bbox=box,words=[dict(text=recognized,bbox=box,confidence=96)]))
    annotate_lines(lines,np.array(image))
    assert all(11<=line['font_size_pt']<=13 for line in lines)
    assert abs(lines[0]['font_size_pt']-lines[1]['font_size_pt'])<=.5


def test_table_keeps_real_small_and_large_type_levels():
    lines=[dict(text='ordinary text',font_size_pt=12,font_name='Times New Roman') for _ in range(8)]
    lines += [dict(text='a smaller note',font_size_pt=8,font_name='Times New Roman'),
              dict(text='LARGER HEADING',font_size_pt=18,font_name='Times New Roman')]
    table=dict(type='table',cells=[dict(text=l['text'],source_lines=[l]) for l in lines])
    harmonize_typography([dict(elements=[table])])
    assert [l['font_size_pt'] for l in lines]==[12]*8+[8,18]


def test_missed_comma_tail_does_not_enlarge_numeric_font():
    image=Image.new('L',(700,100),255);draw=ImageDraw.Draw(image);font=font_at(50,True)
    lines=[]
    for x,printed,recognized in [(10,'28,6','286'),(350,'47,2','47,2')]:
        box=list(draw.textbbox((x,10),printed,font=font));draw.text((x,10),printed,font=font,fill=0)
        lines.append(dict(text=recognized,bbox=box,words=[dict(text=recognized,bbox=box,confidence=96)]))
    annotate_lines(lines,np.array(image))
    assert all(11<=line['font_size_pt']<=13 for line in lines)


def test_missing_stacked_accent_does_not_enlarge_an_alphabetic_word():
    image=Image.new('L',(700,160),255);draw=ImageDraw.Draw(image);font=font_at(50)
    lines=[]
    for x,printed,recognized in [(10,'Nguyễn','Nguyên'),(350,'Thắng','Thăng')]:
        box=list(draw.textbbox((x,10),printed,font=font));draw.text((x,10),printed,font=font,fill=0)
        lines.append(dict(text=recognized,bbox=box,words=[dict(text=recognized,bbox=box,confidence=96)]))
    annotate_lines(lines,np.array(image))
    assert all(11<=line['font_size_pt']<=13 for line in lines)


def test_one_ambiguous_line_does_not_change_paragraph_family_or_size():
    lines=[dict(text='A long ordinary line of body text',font_name='Times New Roman',font_size_pt=12),
           dict(text='A single uncertain line',font_name='Calibri',font_size_pt=15,
                font_fits={'Times New Roman':[.4,False,False,13], 'Calibri':[.1,False,False,15]}),
           dict(text='Another long ordinary line of body text',font_name='Times New Roman',font_size_pt=12)]
    page=dict(elements=[dict(type='paragraph',source_lines=lines)])
    harmonize_typography([page])
    assert {(l['font_name'],l['font_size_pt']) for l in lines}=={('Times New Roman',12)}


def test_export_keeps_glyph_proportions_and_superscript_when_line_is_tight(tmp_path):
    from docx.oxml.ns import qn
    line=dict(text='A compact sentence7',font_name='Times New Roman',font_size_pt=12,
              bbox=[0,0,300,45],inline_runs=[{'text':'A compact sentence'},
                                          {'text':'7','vertical_align':'superscript'}])
    path=tmp_path/'tracking.docx'
    build_docx({'pages':[{'page_w':2480,'elements':[dict(type='paragraph',text=line['text'],source_lines=[line],bbox=line['bbox'])]}]},path)
    runs=Document(path).paragraphs[0].runs
    assert ''.join(r.text for r in runs)==line['text']
    assert all(r._r.find('.//'+qn('w:w')) is None for r in runs)
    assert any(r._r.find('.//'+qn('w:spacing')) is not None for r in runs)
    assert runs[-1].font.superscript


def test_small_table_paragraph_mark_does_not_inherit_body_font_size(tmp_path):
    from docx.oxml.ns import qn
    path=tmp_path/'small.docx'
    line=dict(text='123',bbox=[8,4,35,25],font_name='Times New Roman',font_size_pt=6)
    cell=dict(row=0,col=0,rowspan=1,colspan=1,text='123',source_lines=[line])
    table=dict(type='table',bbox=[0,0,300,40],xs=[0,300],ys=[0,40],grid=[['123']],cells=[cell])
    build_docx(dict(pages=[dict(page_w=2480,elements=[table])]),path)
    p=Document(path).tables[0].cell(0,0).paragraphs[0]
    assert p._p.xpath('./w:pPr/w:rPr/w:sz')[0].get(qn('w:val'))=='12'


def script_image():
    image=Image.new('L',(600,140),255);draw=ImageDraw.Draw(image)
    normal=font_at(48);small=font_at(28)
    words=[]
    for x,text in [(10,'Total'),(150,'value')]:
        b=list(draw.textbbox((x,45),text,font=normal));draw.text((x,45),text,font=normal,fill=0)
        words.append(dict(text=text,bbox=b,confidence=96))
    x=words[-1]['bbox'][2]+5
    b=list(draw.textbbox((x,35),'7',font=small));draw.text((x,35),'7',font=small,fill=0)
    # First detector mistakes the raised glyph for trailing punctuation.
    words[-1]['text']+='”';words[-1]['bbox'][2]=b[2]
    box=[10,min(w['bbox'][1] for w in words+[dict(bbox=b)]),b[2],max(w['bbox'][3] for w in words)]
    return np.array(image),dict(text='Total value”',bbox=box,words=words),b


def test_raised_digit_is_recovered_and_exported_as_superscript(tmp_path):
    image,line,box=script_image()
    assert len(script_candidates(line,image))==1
    recover_inline_scripts([line],image,lambda crops:[('7',.99)],
                           lambda crop,isolated:('7',96) if isolated else ('value',96),lambda:False)
    assert line['text']=='Total value7'
    assert [r['text'] for r in line['inline_runs'] if r.get('vertical_align')]==['7']
    path=tmp_path/'script.docx'
    line.update(font_size_pt=12,font_name='Times New Roman')
    build_docx(dict(pages=[dict(elements=[dict(type='paragraph',text=line['text'],source_lines=[line])])]),path)
    runs=Document(path).paragraphs[0].runs
    assert ''.join(r.text for r in runs)=='Total value7'
    assert [(r.text,r.font.superscript) for r in runs if r.font.superscript]==[('7',True)]


def test_script_disagreement_and_nonnumeric_readings_are_not_replaced():
    for reading,verification in [('7','1'),('l','1')]:
        image,line,_=script_image();original=deepcopy(line)
        recover_inline_scripts([line],image,lambda crops:[(reading,.99)],
                               lambda crop,isolated:(verification,96),lambda:False)
        assert line==original


def test_accents_above_letters_are_not_superscripts():
    image=Image.new('L',(600,120),255);draw=ImageDraw.Draw(image);font=font_at(48)
    text='été élevé côté';box=list(draw.textbbox((10,25),text,font=font));draw.text((10,25),text,font=font,fill=0)
    line=dict(text=text,bbox=box,words=[dict(text=text,bbox=box,confidence=96)])
    assert script_candidates(line,np.array(image))==[]


def test_curved_quotes_are_not_promoted_to_repeated_superscript_digits():
    from app.services.ocr_inline import _raised_symbol
    for bold in [False,True]:
        for quote,digits in [('“','66'),('”','99')]:
            image=Image.new('L',(100,100),255)
            ImageDraw.Draw(image).text((10,10),quote,font=font_at(65,bold),fill=0)
            assert _raised_symbol(np.array(image),digits)==quote
    for digits in ['6','9','66','99','17']:
        image=Image.new('L',(150,100),255)
        ImageDraw.Draw(image).text((10,10),digits,font=font_at(30),fill=0)
        assert _raised_symbol(np.array(image),digits)==digits


def test_visual_accent_choice_is_bidirectional_and_abstains_on_other_letters():
    for printed,wrong in [('côté','côtè'),('côtè','côté')]:
        font=font_at(65,True,True);image=Image.new('L',(300,110),255);draw=ImageDraw.Draw(image)
        draw.text((10,10),printed,font=font,fill=0);crop=np.array(image)
        assert visual_accent_support(crop,wrong,printed)
        assert not visual_accent_support(crop,printed,wrong)
        assert not visual_accent_support(crop,printed,'test')


def test_language_decoder_omission_does_not_delete_small_numbers_or_symbols():
    line=dict(text='A 7 + B',bbox=[10,10,110,40],words=[
        dict(text=t,bbox=[10+i*25,10,30+i*25,40],confidence=15) for i,t in enumerate(['A','7','+','B'])])
    class Reader:
        def predict(self,crop):raise AssertionError('Deletion is not visual evidence')
    refine_weak_spans(line,'A B',np.full((80,150),255,np.uint8),Reader(),lambda:False)
    assert line['text']=='A 7 + B'


def test_confident_missing_letter_needs_independent_crop_agreement():
    original=dict(text='An A example',bbox=[5,5,130,40],words=[
        dict(text=t,bbox=[5+i*40,5,35+i*40,40],confidence=95) for i,t in enumerate(['An','A','example'])])
    class Reader:
        def predict(self,crop):return 'An Ab example'
    for verified,score,expected in [('Ab',92,'An Ab example'),('A',95,'An A example'),('Ab',60,'An A example')]:
        line=deepcopy(original)
        refine_with_context(line,'An Ab example',np.full((60,150),255,np.uint8),Reader(),lambda:False,
                            verify_word=lambda crop:dict(text=verified,confidence=score))
        assert line['text']==expected


def test_raised_digit_outside_detector_top_is_still_considered():
    image,line,_=script_image()
    line['bbox'][1]=min(w['bbox'][1] for w in line['words'])
    assert len(script_candidates(line,image))==1


def test_overlapping_detector_box_does_not_enlarge_font():
    image=Image.new('L',(900,190),255);draw=ImageDraw.Draw(image);font=font_at(50)
    lines=[]
    for y in [10,85]:
        words=[]
        for x,text in [(10,'Printed'),(210,'normal'),(410,'letters')]:
            box=list(draw.textbbox((x,y),text,font=font));draw.text((x,y),text,font=font,fill=0)
            words.append(dict(text=text,bbox=box,confidence=96))
        lines.append(dict(text='Printed normal letters',words=words,
                          bbox=[10,min(w['bbox'][1] for w in words),words[-1]['bbox'][2],max(w['bbox'][3] for w in words)]))
    lines[0]['words'][0]['bbox'][3]=lines[1]['words'][0]['bbox'][3]
    lines[0]['bbox'][3]=lines[1]['bbox'][3]
    annotate_lines(lines,np.array(image))
    assert all(11<=l['font_size_pt']<=13 for l in lines)
    assert abs(lines[0]['font_size_pt']-lines[1]['font_size_pt'])<=.5


def test_right_anchored_cell_uses_alignment_instead_of_large_indent(tmp_path):
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    line=dict(text='Example',bbox=[210,15,294,50],font_name='Times New Roman',font_size_pt=9)
    table=dict(type='table',bbox=[0,0,300,80],xs=[0,300],ys=[0,80],grid=[['Example']],
               cells=[dict(row=0,col=0,rowspan=1,colspan=1,text='Example',source_lines=[line])])
    path=tmp_path/'right.docx'
    build_docx(dict(pages=[dict(page_w=2480,elements=[table])]),path)
    p=Document(path).tables[0].cell(0,0).paragraphs[0]
    assert p.alignment==WD_ALIGN_PARAGRAPH.RIGHT
    assert not p.paragraph_format.left_indent


def test_small_breve_and_circumflex_are_resolved_in_both_directions():
    for size in [24,32,40]:
        for printed,wrong in [('cấn','cắn'),('cắn','cấn')]:
            image=Image.new('L',(140,90),255)
            ImageDraw.Draw(image).text((10,10),printed,font=font_at(size,True),fill=0)
            assert visual_accent_support(np.array(image),wrong,printed)
            assert not visual_accent_support(np.array(image),printed,wrong)


def test_registered_missing_tone_does_not_move_unchanged_letter_bodies():
    from app.services.ocr_glyph_evidence import registered_accent_support
    for family,bold,italic in [('Arial',False,False),('Times New Roman',False,True)]:
        font=next(f for n,b,i,f in _fonts() if (n,b,i)==(family,bold,italic)).font_variant(size=48)
        for printed,wrong in [('tổ','tô'),('dẫn','dân'),('đổi','đôi')]:
            im=Image.new('L',(240,120),255)
            ImageDraw.Draw(im).text((10,10),printed,font=font,fill=0)
            assert registered_accent_support(np.array(im),wrong,printed)
            assert not registered_accent_support(np.array(im),printed,wrong)


def test_period_removal_requires_actual_missing_ink_even_if_readers_omit_it():
    for printed in ['mạng','mạng.']:
        image=Image.new('L',(360,120),255);draw=ImageDraw.Draw(image)
        font=font_at(50);box=list(draw.textbbox((20,10),printed,font=font))
        draw.text((20,10),printed,font=font,fill=0)
        line=dict(text='mạng.',words=[dict(text='mạng.',bbox=box,confidence=95)])
        class Reader:
            def predict(self,crop):return 'mạng'
        refine_with_context(line,'mạng',np.array(image),Reader(),lambda:False,
                            verify_word=lambda c:dict(text='mạng',confidence=96))
        assert line['text']==printed


def test_weak_weight_outlier_is_regular_but_strong_emphasis_and_heading_survive():
    def line(text,bold=False,margin=.03):
        return dict(text=text,font_name='Arial',font_size_pt=12,bold=bold,
                    font_fits={'Arial':[.18,bold,False,12]},
                    font_bold_margins={'Arial':{'roman':margin}})
    body=[line('A long ordinary line of printed body text.') for _ in range(4)]
    weak=line('The last short line.',True)
    strong=line('A strongly emphasized line.',True,.13)
    heading=line('A leading heading',True)
    pages=[dict(elements=[dict(type='paragraph',source_lines=[heading]+body+[weak]),
                          dict(type='paragraph',source_lines=body+[strong])])]
    harmonize_typography(pages)
    assert not weak['bold']
    assert strong['bold'] and heading['bold']
