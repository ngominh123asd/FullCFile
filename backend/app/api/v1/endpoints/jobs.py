"""
FileForge Backend - Jobs Endpoints - No auth required
"""
import os
from pathlib import Path
from typing import Optional
from fastapi import APIRouter, Depends, HTTPException, status, Query
from fastapi.responses import FileResponse
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy import select

from app.core.database import get_db
from app.api.v1.endpoints.files import ensure_guest_user
from app.models.database import Job, OutputFile, File

router = APIRouter()


@router.get("/{job_id}")
async def get_job(
    job_id: str,
    db: AsyncSession = Depends(get_db),
):
    """Get job status"""
    user_id = await ensure_guest_user(db)

    result = await db.execute(
        select(Job).where(
            Job.id == job_id,
            Job.user_id == user_id,
        )
    )
    job = result.scalar_one_or_none()

    if not job:
        raise HTTPException(status_code=404, detail="Job not found")

    input_filename = ""
    if job.input_file_id:
        f_res = await db.execute(select(File).where(File.id == job.input_file_id))
        f_obj = f_res.scalar_one_or_none()
        if f_obj:
            input_filename = f_obj.original_filename or f_obj.filename

    return {
        "id": job.id,
        "user_id": job.user_id,
        "input_file_id": job.input_file_id,
        "input_filename": input_filename,
        "job_type": job.job_type,
        "status": job.status,
        "output_format": job.output_format,
        "progress": job.progress,
        "stage": job.stage,
        "error": job.error,
        "result": job.result,
        "created_at": job.created_at,
        "started_at": job.started_at,
        "completed_at": job.completed_at,
    }


@router.get("/{job_id}/download")
async def download_job_result(
    job_id: str,
    db: AsyncSession = Depends(get_db),
):
    """Download job result file"""
    user_id = await ensure_guest_user(db)

    result = await db.execute(
        select(Job).where(
            Job.id == job_id,
            Job.user_id == user_id,
        )
    )
    job = result.scalar_one_or_none()

    if not job:
        raise HTTPException(status_code=404, detail="Job not found")

    if job.status != "completed":
        raise HTTPException(
            status_code=400,
            detail=f"Job is not completed. Current status: {job.status}",
        )

    if not job.result or "output_path" not in job.result:
        raise HTTPException(status_code=404, detail="Output file not found")

    output_path = job.result["output_path"]

    if not os.path.exists(output_path):
        raise HTTPException(status_code=404, detail="Output file not found on storage")

    media_types = {
        "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "pdf": "application/pdf",
        "html": "text/html",
        "txt": "text/plain",
        "md": "text/markdown",
        "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "csv": "text/csv",
        "json": "application/json",
    }

    media_type = media_types.get(job.output_format, "application/octet-stream")
    
    original_stem = f"fileforge_{job.id[:8]}"
    if job.input_file_id:
        file_res = await db.execute(select(File).where(File.id == job.input_file_id))
        input_f = file_res.scalar_one_or_none()
        if input_f:
            fn = input_f.original_filename or input_f.filename
            if fn:
                original_stem = Path(fn).stem

    filename = f"{original_stem}.{job.output_format}"

    return FileResponse(
        path=output_path,
        filename=filename,
        media_type=media_type,
    )


@router.get("/{job_id}/preview")
async def preview_job_result(
    job_id: str,
    db: AsyncSession = Depends(get_db),
):
    """Preview job result - returns content for inline display"""
    user_id = await ensure_guest_user(db)

    result = await db.execute(
        select(Job).where(
            Job.id == job_id,
            Job.user_id == user_id,
        )
    )
    job = result.scalar_one_or_none()

    if not job:
        raise HTTPException(status_code=404, detail="Job not found")

    if job.status != "completed":
        raise HTTPException(status_code=400, detail="Job is not completed")

    if not job.result or "output_path" not in job.result:
        raise HTTPException(status_code=404, detail="Output file not found")

    output_path = job.result["output_path"]

    if not os.path.exists(output_path):
        raise HTTPException(status_code=404, detail="Output file not found on storage")

    dur_seconds = None
    if job.started_at and job.completed_at:
        dur_seconds = round((job.completed_at - job.started_at).total_seconds(), 1)
    elif job.result and isinstance(job.result, dict) and "duration_seconds" in job.result:
        dur_seconds = job.result["duration_seconds"]

    # Read content for preview-able formats
    if job.output_format in ["txt", "md", "html", "csv", "json"]:
        try:
            with open(output_path, "r", encoding="utf-8") as f:
                text = f.read()
            return {
                "job_id": job.id,
                "format": job.output_format,
                "content": text,
                "previewable": True,
                "started_at": job.started_at.isoformat() if job.started_at else None,
                "completed_at": job.completed_at.isoformat() if job.completed_at else None,
                "duration_seconds": dur_seconds,
            }
        except Exception as e:
            raise HTTPException(status_code=500, detail=f"Failed to read: {str(e)}")

    # Binary formats
    elif job.output_format == "docx":
        try:
            from docx import Document
            from docx.oxml.ns import qn
            import re
            import html as html_module

            doc = Document(output_path)

            NS_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"

            def is_xml_attr_true(parent_elem, tag_name):
                """Check if boolean XML element (like w:b, w:i) is genuinely true in Word OpenXML."""
                if parent_elem is None:
                    return False
                elem = parent_elem.find(NS_W + tag_name)
                if elem is None:
                    return False
                val = elem.get(NS_W + "val")
                if val is None:
                    return True
                val_lower = val.strip().lower()
                return val_lower in ("1", "true", "on")

            def is_underline_true(parent_elem):
                if parent_elem is None:
                    return False
                elem = parent_elem.find(NS_W + "u")
                if elem is None:
                    return False
                val = elem.get(NS_W + "val")
                if val is None:
                    return True
                val_lower = val.strip().lower()
                return val_lower not in ("none", "0", "false", "off")

            from app.services.document_preview import page_setup, safe_float, safe_int
            section_setups = [page_setup(section) for section in doc.sections]
            if not section_setups:
                section_setups = [{
                    'width_pt': 595.3, 'height_pt': 841.9,
                    'top_pt': 56.7, 'bottom_pt': 56.7,
                    'left_pt': 56.7, 'right_pt': 56.7
                }]
            source_pages = 'fileforge:source-pages' in (doc.core_properties.keywords or '').split(',')

            normal_line_spacing = 1.15
            normal_font_name = "Times New Roman"
            normal_font_size_pt = 11.0
            try:
                normal = doc.styles['Normal']
                if normal.paragraph_format and normal.paragraph_format.line_spacing:
                    sp_val = normal.paragraph_format.line_spacing
                    if isinstance(sp_val, (int, float)):
                        normal_line_spacing = float(sp_val)
                if normal.font:
                    if normal.font.name:
                        normal_font_name = normal.font.name
                    if normal.font.size and hasattr(normal.font.size, 'pt'):
                        normal_font_size_pt = float(normal.font.size.pt)
            except Exception:
                pass

            def parse_runs_html_and_props(p_elem, is_table_cell: bool = False):
                """Extract formatted HTML, style, font, size, alignment, and page break info."""
                runs_html = []
                image_height = 0.0
                style_name = "Normal"
                align = "left"
                has_page_break = False
                margin_top_pt = 0.0
                margin_bottom_pt = 0.0
                indent_left_pt = indent_right_pt = first_line_pt = 0.0
                line_height_pt = None
                line_spacing = normal_line_spacing

                pPr = p_elem.find(NS_W + "pPr")
                if pPr is not None:
                    pStyle = pPr.find(NS_W + "pStyle")
                    if pStyle is not None:
                        style_name = pStyle.get(NS_W + "val") or "Normal"

                    jc = pPr.find(NS_W + "jc")
                    if jc is not None:
                        jc_val = (jc.get(NS_W + "val") or "left").lower()
                        if jc_val in ("center", "right", "both", "distribute", "justify"):
                            align = "justify" if jc_val in ("both", "distribute", "justify") else jc_val

                    sp = pPr.find(NS_W + "spacing")
                    if sp is not None:
                        b_val = sp.get(NS_W + "before")
                        a_val = sp.get(NS_W + "after")
                        if b_val:
                            margin_top_pt = safe_float(b_val) / 20.0
                        if a_val:
                            margin_bottom_pt = safe_float(a_val) / 20.0
                        line = sp.get(NS_W + 'line')
                        if line:
                            line_val = safe_float(line)
                            if sp.get(NS_W + 'lineRule', 'auto') == 'auto':
                                line_spacing = line_val / 240.0
                            else:
                                line_height_pt = line_val / 20.0
                    indent = pPr.find(NS_W + 'ind')
                    if indent is not None:
                        indent_left_pt = safe_float(indent.get(NS_W + 'left', '0')) / 20.0
                        indent_right_pt = safe_float(indent.get(NS_W + 'right', '0')) / 20.0
                        first_line_pt = (
                            safe_float(indent.get(NS_W + 'firstLine', '0'))
                            - safe_float(indent.get(NS_W + 'hanging', '0'))
                        ) / 20.0

                    pgb = pPr.find(NS_W + "pageBreakBefore")
                    if pgb is not None:
                        has_page_break = True

                primary_font = normal_font_name
                primary_size_pt = normal_font_size_pt
                style_lower = style_name.lower()
                if "heading 1" in style_lower or "title" in style_lower:
                    primary_size_pt = 20.0
                elif "heading 2" in style_lower:
                    primary_size_pt = 16.0
                elif "heading 3" in style_lower:
                    primary_size_pt = 14.0
                elif "heading 4" in style_lower:
                    primary_size_pt = 13.0

                for r in p_elem.iter(NS_W + "r"):
                    rPr = r.find(NS_W + "rPr")
                    rfont = ""
                    rsize = primary_size_pt
                    bold = False
                    italic = False
                    underline = False
                    is_sup = is_sub = False
                    color = ""

                    if rPr is not None:
                        rFonts = rPr.find(NS_W + "rFonts")
                        if rFonts is not None:
                            rfont = rFonts.get(NS_W + "ascii") or rFonts.get(NS_W + "hAnsi") or ""
                            if not primary_font:
                                primary_font = rfont

                        sz = rPr.find(NS_W + "sz")
                        if sz is not None and sz.get(NS_W + "val"):
                            rsize = safe_float(sz.get(NS_W + "val")) / 2.0
                            primary_size_pt = rsize

                        bold = is_xml_attr_true(rPr, "b")
                        italic = is_xml_attr_true(rPr, "i")
                        underline = is_underline_true(rPr)
                        vertAlign_el = rPr.find(NS_W + "vertAlign")
                        is_sup = vertAlign_el is not None and vertAlign_el.get(NS_W + "val") == "superscript"
                        is_sub = vertAlign_el is not None and vertAlign_el.get(NS_W + "val") == "subscript" 

                        color_el = rPr.find(NS_W + "color")
                        if color_el is not None and color_el.get(NS_W + "val"):
                            color = color_el.get(NS_W + "val")

                    # Check for explicit page break in run
                    for br in r.findall(NS_W + "br"):
                        if br.get(NS_W + "type") == "page":
                            has_page_break = True
                    if r.find(NS_W + "lastRenderedPageBreak") is not None:
                        has_page_break = True

                    from app.services.document_preview import run_text, run_images
                    images, height = run_images(r, doc.part)
                    if images:
                        runs_html.append(images)
                        image_height = max(image_height, height)
                    text = run_text(r)
                    if not text:
                        continue

                    # Replace Private Use Area bullet symbols with standard clean bullets
                    pua_bullets = {
                        '\uf0b7': '• ', '\uf0a7': '▪ ', '\uf0d8': '➢ ', '\uf076': '❖ ',
                        '\uf06e': '■ ', '\uf0fc': '✔ ', '\uf0a8': '◻ ', '\uf0de': '⇒ ',
                        '\uf0e0': '✉ ', '\uf034': '✓ ', '\uf075': '◆ ', '\uf02d': '– ',
                    }
                    for pua_char, rep in pua_bullets.items():
                        if pua_char in text:
                            text = text.replace(pua_char, rep)
                    text = re.sub(r'[\uf000-\uf8ff]', '• ', text)

                    safe = html_module.escape(text)
                    safe = safe.replace("\t", "&emsp;").replace("\n", "<br/>")

                    is_code = "courier" in rfont.lower() if rfont else False
                    if is_code:
                        safe = f"<code>{safe}</code>"
                    if bold:
                        safe = f"<strong>{safe}</strong>"
                    if italic:
                        safe = f"<em>{safe}</em>"
                    if underline:
                        safe = f"<u>{safe}</u>"
                    if is_sup:
                        safe = f"<sup>{safe}</sup>"
                    elif is_sub:
                        safe = f"<sub>{safe}</sub>"
                    stretch = rPr.find(NS_W + 'w') if rPr is not None else None
                    stretch = safe_float(stretch.get(NS_W + 'val', '100')) / 100.0 if stretch is not None else 1.0
                    font_css = f'font-size:{rsize}pt;font-family:{html_module.escape(rfont or primary_font, quote=True)}'
                    tracking = rPr.find(NS_W + 'spacing') if rPr is not None else None
                    if tracking is not None:
                        font_css += f';letter-spacing:{safe_float(tracking.get(NS_W + "val", "0")) / 20.0}pt'
                    if stretch != 1.0 and not is_table_cell:
                        try:
                            from app.services.ocr_typography import text_width_pt
                            parts = []
                            for plain in text.split('\n'):
                                part = html_module.escape(plain).replace('\t', '&emsp;')
                                if bold: part = f'<strong>{part}</strong>'
                                if italic: part = f'<em>{part}</em>'
                                if underline: part = f'<u>{part}</u>'
                                if is_sup: part = f'<sup>{part}</sup>'
                                elif is_sub: part = f'<sub>{part}</sub>'
                                width = text_width_pt(plain, rfont or primary_font, rsize, bold, italic)
                                if width:
                                    part = (f'<span style="display:inline-block;text-indent:0;width:{width*stretch:.3f}pt;{font_css}">'
                                            f'<span style="display:inline-block;text-indent:0;white-space:pre;transform-origin:left;transform:scaleX({stretch})">{part}</span></span>')
                                parts.append(part)
                            safe = '<br/>'.join(parts)
                        except Exception:
                            safe = f'<span style="{font_css}">{safe}</span>'
                    else:
                        safe = f'<span style="{font_css}">{safe}</span>'

                    runs_html.append(safe)

                content = "".join(runs_html)

                # Accurate plain text length measurement
                clean_text = re.sub(r'<[^>]+>', '', content).replace("&emsp;", " ").replace("&nbsp;", " ").strip()
                plain_len = len(clean_text)

                if not plain_len:
                    estimated_height = 14.0
                else:
                    explicit_lines = max(1, content.count("<br/>") + 1)
                    wrap_lines = max(1, int(plain_len / 85.0) + (1 if plain_len % 85 > 0 else 0))
                    lines_count = max(explicit_lines, wrap_lines)
                    line_height = primary_size_pt * 1.25

                    if "heading 1" in style_lower or "title" in style_lower:
                        spacing_extra = 12.0
                    elif "heading 2" in style_lower:
                        spacing_extra = 10.0
                    elif "heading 3" in style_lower:
                        spacing_extra = 8.0
                    elif "heading 4" in style_lower:
                        spacing_extra = 6.0
                    else:
                        spacing_extra = 2.0

                    estimated_height = (lines_count * line_height) + spacing_extra

                return {
                    "type": "paragraph",
                    "style": style_name,
                    "content": content,
                    "font_name": primary_font or "Times New Roman",
                    "font_size": primary_size_pt,
                    "align": align,
                    "margin_top_pt": margin_top_pt,
                    "margin_bottom_pt": margin_bottom_pt,
                    "indent_left_pt": indent_left_pt,
                    "indent_right_pt": indent_right_pt,
                    "first_line_pt": first_line_pt,
                    "line_height_pt": line_height_pt,
                    "line_spacing": line_spacing,
                    "has_page_break": has_page_break,
                    "height": max(estimated_height, image_height),
                }

            def parse_table_elem(tbl_elem):
                from app.services.document_preview import parse_table
                styles_elem = getattr(doc.styles, 'element', None)
                return parse_table(tbl_elem, parse_runs_html_and_props, styles_elem)

            # Step 1: Collect all block items in sequential body order
            body_blocks = []
            body_elm = doc.element.body
            section_index = 0
            for child in body_elm.iterchildren():
                tag = child.tag
                if tag.endswith("}p"):
                    p_info = parse_runs_html_and_props(child)
                    p_info['_page_setup'] = section_setups[min(section_index, len(section_setups) - 1)]
                    sect = child.find('./' + NS_W + 'pPr/' + NS_W + 'sectPr')
                    if sect is not None:
                        next_idx = min(section_index + 1, len(doc.sections) - 1)
                        next_section = doc.sections[next_idx] if next_idx < len(doc.sections) else None
                        is_new_page = False
                        if next_section is not None:
                            try:
                                is_new_page = safe_int(getattr(next_section, 'start_type', 0)) != 0
                            except Exception:
                                is_new_page = False
                        p_info['section_end'] = source_pages or is_new_page
                        section_index += 1
                    body_blocks.append(p_info)
                elif tag.endswith("}tbl"):
                    t_info = parse_table_elem(child)
                    if t_info["data"]:
                        t_info['_page_setup'] = section_setups[min(section_index, len(section_setups) - 1)]
                        body_blocks.append(t_info)
                elif tag.endswith("}sdt"):
                    sdt_content = child.find('./' + NS_W + 'sdtContent')
                    if sdt_content is not None:
                        for sdt_child in sdt_content.iterchildren():
                            sdt_tag = sdt_child.tag
                            if sdt_tag.endswith("}p"):
                                p_info = parse_runs_html_and_props(sdt_child)
                                p_info['_page_setup'] = section_setups[min(section_index, len(section_setups) - 1)]
                                body_blocks.append(p_info)
                            elif sdt_tag.endswith("}tbl"):
                                t_info = parse_table_elem(sdt_child)
                                if t_info["data"]:
                                    t_info['_page_setup'] = section_setups[min(section_index, len(section_setups) - 1)]
                                    body_blocks.append(t_info)

            # Step 2: Paginate sequentially into clean pages
            pages = []
            page_setups = []
            current_page_blocks = []
            current_page_height = 0.0
            current_setup = section_setups[0]

            def finish_page(force=False):
                nonlocal current_page_blocks, current_page_height
                if current_page_blocks or force:
                    pages.append(current_page_blocks)
                    page_setups.append(current_setup)
                current_page_blocks = []
                current_page_height = 0.0

            for block in body_blocks:
                is_explicit_break = block.get("has_page_break", False)

                if is_explicit_break and current_page_blocks:
                    finish_page()
                    # If this was an empty page-break marker paragraph, don't put an empty spacer at top of next page
                    if block["type"] == "paragraph" and not block["content"].strip():
                        continue

                current_setup = block.pop('_page_setup', current_setup)
                usable_height = current_setup['height_pt'] - current_setup['top_pt'] - current_setup['bottom_pt']
                if block.get('section_end'):
                    if block.get('content', '').strip():
                        current_page_blocks.append(block)
                    finish_page(force=source_pages)
                    continue

                block_h = block.get("height", 18.0)

                # Split at row boundaries while keeping vertical merges intact.
                if not source_pages and block["type"] == "table" and block_h > usable_height:
                    from app.services.document_preview import table_chunks
                    if current_page_blocks:
                        finish_page()
                    chunks = list(table_chunks(block, usable_height))
                    for chunk in chunks[:-1]:
                        pages.append([chunk])
                        page_setups.append(current_setup)
                    if chunks:
                        current_page_blocks = [chunks[-1]]
                        current_page_height = chunks[-1]['height']
                    continue

                if not source_pages and current_page_height + block_h > usable_height and current_page_height > 100.0:
                    finish_page()
                    current_page_blocks = [block]
                    current_page_height = block_h
                else:
                    current_page_blocks.append(block)
                    current_page_height += block_h

            finish_page(force=source_pages)

            if not pages:
                pages = [[]]
                page_setups = [section_setups[0]]

            # Flatten legacy paragraphs & tables for full backward compatibility
            flat_paragraphs = []
            flat_tables = []
            for page in pages:
                for b in page:
                    if b["type"] == "paragraph":
                        flat_paragraphs.append({
                            "style": b["style"],
                            "content": b["content"],
                            "font_name": b.get("font_name", "Times New Roman"),
                            "align": b.get("align", "left"),
                        })
                    elif b["type"] == "table":
                        flat_tables.append(b["data"])
                flat_paragraphs.append({"style": "_page_break", "content": ""})

            return {
                "job_id": job.id,
                "format": "docx",
                "pages": pages,
                "page_count": len(pages),
                "page_setups": page_setups,
                "paragraphs": flat_paragraphs,
                "tables": flat_tables,
                "previewable": True,
                "started_at": job.started_at.isoformat() if job.started_at else None,
                "completed_at": job.completed_at.isoformat() if job.completed_at else None,
                "duration_seconds": dur_seconds,
            }
        except Exception as e:
            # Fallback docx preview: extract basic text & tables if advanced AST parsing encounters unexpected issues
            try:
                from docx import Document
                doc = Document(output_path)
                fallback_blocks = []
                for p in doc.paragraphs:
                    text = p.text
                    if text:
                        fallback_blocks.append({
                            "type": "paragraph",
                            "style": p.style.name if p.style else "Normal",
                            "content": html_module.escape(text),
                            "font_name": "Times New Roman",
                            "font_size": 12.0,
                            "align": "left",
                        })
                for tbl in doc.tables:
                    tbl_data = []
                    for row in tbl.rows:
                        tbl_data.append([html_module.escape(cell.text) for cell in row.cells])
                    if tbl_data:
                        fallback_blocks.append({
                            "type": "table",
                            "data": tbl_data,
                            "has_borders": True,
                        })
                if fallback_blocks:
                    return {
                        "job_id": job.id,
                        "format": "docx",
                        "pages": [fallback_blocks],
                        "page_count": 1,
                        "page_setups": [{'width_pt': 595.3, 'height_pt': 841.9, 'top_pt': 56.7, 'bottom_pt': 56.7, 'left_pt': 56.7, 'right_pt': 56.7}],
                        "paragraphs": fallback_blocks,
                        "tables": [],
                        "previewable": True,
                        "started_at": job.started_at.isoformat() if job.started_at else None,
                        "completed_at": job.completed_at.isoformat() if job.completed_at else None,
                        "duration_seconds": dur_seconds,
                    }
            except Exception:
                pass

            return {
                "job_id": job.id,
                "format": "docx",
                "error": str(e),
                "previewable": False,
                "started_at": job.started_at.isoformat() if job.started_at else None,
                "completed_at": job.completed_at.isoformat() if job.completed_at else None,
                "duration_seconds": dur_seconds,
            }

    elif job.output_format == "xlsx":
        try:
            import openpyxl
            wb = openpyxl.load_workbook(output_path, data_only=True)
            sheets = []
            for sheet_name in wb.sheetnames:
                ws = wb[sheet_name]
                data, warnings = [], {}
                for row in ws.iter_rows(max_row=min(ws.max_row, 100)):
                    if any(cell.value is not None for cell in row):
                        ri = len(data)
                        data.append([str(c.value) if c.value is not None else "" for c in row])
                        for ci, cell in enumerate(row):
                            if cell.comment and cell.comment.author == 'FileForge OCR':
                                warnings[f'{ri}:{ci}'] = cell.comment.text
                sheets.append({"name": sheet_name, "data": data, "cell_warnings": warnings})


            return {
                "job_id": job.id,
                "format": "xlsx",
                "sheets": sheets,
                "previewable": True,
                "started_at": job.started_at.isoformat() if job.started_at else None,
                "completed_at": job.completed_at.isoformat() if job.completed_at else None,
                "duration_seconds": dur_seconds,
            }
        except Exception as e:
            return {
                "job_id": job.id,
                "format": "xlsx",
                "error": str(e),
                "previewable": False,
                "started_at": job.started_at.isoformat() if job.started_at else None,
                "completed_at": job.completed_at.isoformat() if job.completed_at else None,
                "duration_seconds": dur_seconds,
            }

    elif job.output_format == "pdf":
        # Return file info for PDF.js to render
        file_size = os.path.getsize(output_path)
        return {
            "job_id": job.id,
            "format": "pdf",
            "file_size": file_size,
            "url": f"/api/v1/jobs/{job_id}/download",
            "previewable": True,
            "render_type": "pdf",
            "started_at": job.started_at.isoformat() if job.started_at else None,
            "completed_at": job.completed_at.isoformat() if job.completed_at else None,
            "duration_seconds": dur_seconds,
        }

    return {
        "job_id": job.id,
        "format": job.output_format,
        "previewable": False,
        "started_at": job.started_at.isoformat() if job.started_at else None,
        "completed_at": job.completed_at.isoformat() if job.completed_at else None,
        "duration_seconds": dur_seconds,
    }


@router.get("")
async def list_jobs(
    status_filter: str = Query(None, alias="status"),
    job_type: str = Query(None),
    page: int = Query(1, ge=1),
    limit: Optional[int] = Query(None, ge=1),
    db: AsyncSession = Depends(get_db),
):
    """List jobs. If `limit` is omitted, the full (unlimited) history is returned."""
    user_id = await ensure_guest_user(db)

    query = select(Job).where(Job.user_id == user_id)

    if status_filter:
        query = query.where(Job.status == status_filter)
    if job_type:
        query = query.where(Job.job_type == job_type)

    query = query.order_by(Job.created_at.desc())
    if limit is not None:
        query = query.offset((page - 1) * limit).limit(limit)

    result = await db.execute(query)
    jobs = result.scalars().all()

    input_file_ids = [j.input_file_id for j in jobs if j.input_file_id]
    input_filenames = {}
    if input_file_ids:
        file_res = await db.execute(select(File).where(File.id.in_(input_file_ids)))
        for f in file_res.scalars().all():
            input_filenames[f.id] = f.original_filename or f.filename

    return [
        {
            "id": job.id,
            "input_file_id": job.input_file_id,
            "input_filename": input_filenames.get(job.input_file_id, ""),
            "job_type": job.job_type,
            "status": job.status,
            "output_format": job.output_format,
            "progress": job.progress,
            "stage": job.stage,
            "error": job.error,
            "created_at": job.created_at,
            "completed_at": job.completed_at,
        }
        for job in jobs
    ]


@router.delete("")
async def clear_all_jobs(
    db: AsyncSession = Depends(get_db),
):
    """Delete all jobs for the guest/current user"""
    user_id = await ensure_guest_user(db)
    from sqlalchemy import delete

    # Fetch all jobs to clean up output files
    result = await db.execute(select(Job).where(Job.user_id == user_id))
    jobs = result.scalars().all()
    for job in jobs:
        if job.result and isinstance(job.result, dict) and "output_path" in job.result:
            output_path = job.result["output_path"]
            if os.path.exists(output_path):
                try:
                    os.remove(output_path)
                except Exception:
                    pass

    await db.execute(delete(Job).where(Job.user_id == user_id))
    return {"message": "All jobs deleted successfully"}


@router.delete("/{job_id}")
async def cancel_job(
    job_id: str,
    db: AsyncSession = Depends(get_db),
):
    """Cancel/delete single job"""
    user_id = await ensure_guest_user(db)

    result = await db.execute(
        select(Job).where(
            Job.id == job_id,
            Job.user_id == user_id,
        )
    )
    job = result.scalar_one_or_none()

    if not job:
        raise HTTPException(status_code=404, detail="Job not found")

    # Clean up output file on disk if exists
    if job.result and isinstance(job.result, dict) and "output_path" in job.result:
        output_path = job.result["output_path"]
        if os.path.exists(output_path):
            try:
                os.remove(output_path)
            except Exception:
                pass

    if job.status in ["pending", "processing"]:
        job.status = "failed"
        job.error = "Job cancelled by user"
        try:
            from app.services.ocr_service import cancel_ocr_job
            cancel_ocr_job(job.id)
        except Exception:
            pass
        await db.flush()

    await db.delete(job)
    return {"message": "Job deleted successfully"}
